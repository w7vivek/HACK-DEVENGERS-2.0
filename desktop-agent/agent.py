import os
import sys
import time
import json
import threading
import subprocess
from datetime import datetime
from pathlib import Path
import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image, ImageDraw

TOKEN_PATH = Path.home() / '.nami_token.json'
EXCEL_FILENAME = 'NAMI_Transactions.xlsx'
API_BASE = os.environ.get('NAMI_API_URL', 'http://localhost:5001/api')
POLL_INTERVAL_SECONDS = 15 * 60

def load_stored_token():
    if TOKEN_PATH.exists():
        try:
            with open(TOKEN_PATH, 'r', encoding='utf-8') as f:
                data = json.load(f)
                return data.get('token')
        except Exception:
            return None
    return None

def save_stored_token(token):
    try:
        with open(TOKEN_PATH, 'w', encoding='utf-8') as f:
            json.dump({'token': token, 'savedAt': datetime.now().isoformat()}, f)
    except Exception as e:
        print(f"Failed to cache token: {e}")

def remove_stored_token():
    if TOKEN_PATH.exists():
        try:
            TOKEN_PATH.unlink()
        except OSError:
            pass

def prompt_login_dialog():
    import tkinter as tk
    from tkinter import messagebox

    result = {'token': None}
    root = tk.Tk()
    root.title("NAMI Desktop Login")
    root.geometry("360x260")
    root.resizable(False, False)
    root.configure(bg="#F7F1E9")

    tk.Label(root, text="NAMI Desktop Sync", font=("Segoe UI", 14, "bold"), bg="#F7F1E9", fg="#1A1A1A").pack(pady=(16, 4))
    tk.Label(root, text="Sign in to sync transactions to Excel", font=("Segoe UI", 9), bg="#F7F1E9", fg="#6B6B6B").pack(pady=(0, 12))

    frame = tk.Frame(root, bg="#F7F1E9")
    frame.pack(padx=24, fill="x")

    tk.Label(frame, text="Email", font=("Segoe UI", 9, "bold"), bg="#F7F1E9", fg="#1A1A1A").pack(anchor="w")
    email_entry = tk.Entry(frame, font=("Segoe UI", 10), relief="flat", highlightthickness=1, highlightbackground="#D1D5DB")
    email_entry.pack(fill="x", pady=(2, 8), ipady=3)

    tk.Label(frame, text="Password", font=("Segoe UI", 9, "bold"), bg="#F7F1E9", fg="#1A1A1A").pack(anchor="w")
    pw_entry = tk.Entry(frame, font=("Segoe UI", 10), show="*", relief="flat", highlightthickness=1, highlightbackground="#D1D5DB")
    pw_entry.pack(fill="x", pady=(2, 12), ipady=3)

    status_label = tk.Label(root, text="", font=("Segoe UI", 8), bg="#F7F1E9", fg="#DC2626")
    status_label.pack()

    def handle_submit():
        email = email_entry.get().strip()
        pw = pw_entry.get().strip()
        if not email or not pw:
            status_label.config(text="Please fill in both email and password")
            return

        status_label.config(text="Verifying credentials...", fg="#F97316")
        root.update()

        try:
            res = requests.post(
                f"{API_BASE}/auth/login",
                json={'email': email, 'password': pw},
                timeout=10
            )
            if res.status_code == 200:
                data = res.json()
                result['token'] = data.get('token')
                save_stored_token(result['token'])
                root.destroy()
            else:
                msg = res.json().get('message', 'Login failed')
                status_label.config(text=msg, fg="#DC2626")
        except Exception as err:
            status_label.config(text=f"Connection error: {err}", fg="#DC2626")

    btn = tk.Button(
        root,
        text="Connect Sync",
        command=handle_submit,
        bg="#F97316",
        fg="white",
        font=("Segoe UI", 10, "bold"),
        relief="flat",
        cursor="hand2",
        padx=16,
        pady=4
    )
    btn.pack(pady=8)

    root.mainloop()
    return result['token']

def ensure_excel_workbook(file_path):
    parent = os.path.dirname(file_path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    if not os.path.exists(file_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Transactions"
        
        headers = [
            "Transaction ID",
            "Timestamp",
            "Sender",
            "Receiver",
            "Amount (INR)",
            "Payment Mode",
            "Account Holder",
            "Account Number",
            "IFSC Code"
        ]
        ws.append(headers)

        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1A1A1A", end_color="1A1A1A", fill_type="solid")
        alignment = Alignment(horizontal="center", vertical="center")

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment

        wb.save(file_path)
    return file_path

def append_transactions_to_excel(file_path, tx_list):
    ensure_excel_workbook(file_path)
    wb = load_workbook(file_path)
    ws = wb["Transactions"] if "Transactions" in wb.sheetnames else wb.active

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    data_font = Font(name="Segoe UI", size=10)
    num_alignment = Alignment(horizontal="right", vertical="center")
    center_alignment = Alignment(horizontal="center", vertical="center")
    left_alignment = Alignment(horizontal="left", vertical="center")

    for tx in tx_list:
        created_str = tx.get('createdAt', '')
        try:
            dt = datetime.fromisoformat(created_str.replace('Z', '+00:00'))
            created_fmt = dt.strftime('%Y-%m-%d %H:%M:%S')
        except Exception:
            created_fmt = created_str

        row = [
            str(tx.get('_id', '')),
            created_fmt,
            tx.get('sender', ''),
            tx.get('receiver', ''),
            float(tx.get('amount', 0)),
            tx.get('mode', ''),
            tx.get('accountHolder', '') or '-',
            tx.get('accountNumber', '') or '-',
            tx.get('ifsc', '') or '-'
        ]
        ws.append(row)

        curr_row = ws.max_row
        for col_idx in range(1, len(row) + 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx == 5:
                cell.alignment = num_alignment
                cell.number_format = '#,##0.00'
            elif col_idx in [1, 2, 6]:
                cell.alignment = center_alignment
            else:
                cell.alignment = left_alignment

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(file_path)

def open_file(file_path):
    try:
        if sys.platform.startswith('win'):
            os.startfile(file_path)
        elif sys.platform == 'darwin':
            subprocess.Popen(['open', file_path])
        else:
            subprocess.Popen(['xdg-open', file_path])
    except Exception as e:
        print(f"Could not open spreadsheet: {e}")

def create_tray_icon():
    width, height = 64, 64
    image = Image.new('RGBA', (width, height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.ellipse([4, 4, width - 4, height - 4], fill="#F97316")
    draw.ellipse([16, 16, width - 16, height - 16], outline="white", width=4)
    draw.text((26, 20), "N", fill="white", font=None)
    return image

class NamiSyncAgent:
    def __init__(self):
        self.token = load_stored_token()
        self.excel_path = os.path.abspath(EXCEL_FILENAME)
        self.is_running = True
        self.tray_icon = None

    def get_headers(self):
        return {
            'Authorization': f"Bearer {self.token}",
            'Content-Type': 'application/json'
        }

    def fetch_unsynced(self):
        if not self.token:
            return []
        try:
            res = requests.get(
                f"{API_BASE}/transactions/unsynced",
                headers=self.get_headers(),
                timeout=10
            )
            if res.status_code == 401:
                remove_stored_token()
                self.token = None
                return []
            if res.status_code == 200:
                return res.json().get('transactions', [])
        except Exception as e:
            print(f"Polling error: {e}")
        return []

    def mark_synced(self, tx_ids):
        if not self.token or not tx_ids:
            return False
        try:
            res = requests.post(
                f"{API_BASE}/transactions/mark-synced",
                headers=self.get_headers(),
                json={'ids': tx_ids},
                timeout=10
            )
            return res.status_code == 200
        except Exception:
            return False

    def prompt_and_sync(self, unsynced_items):
        count = len(unsynced_items)
        if count == 0:
            return

        import tkinter as tk
        from tkinter import messagebox

        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)

        question = f"You have {count} unsaved {'entry' if count == 1 else 'entries'}.\nAdd them to Excel now?"
        confirmed = messagebox.askyesno("NAMI Sync Alert", question)
        root.destroy()

        if confirmed:
            append_transactions_to_excel(self.excel_path, unsynced_items)
            open_file(self.excel_path)
            ids = [t['_id'] for t in unsynced_items if '_id' in t]
            self.mark_synced(ids)
            print(f"Synced {count} entries to {self.excel_path}")

    def run_sync_cycle(self):
        if not self.token:
            self.token = prompt_login_dialog()
            if not self.token:
                return

        unsynced = self.fetch_unsynced()
        if unsynced:
            self.prompt_and_sync(unsynced)

    def background_poller(self):
        while self.is_running:
            try:
                self.run_sync_cycle()
            except Exception as e:
                print(f"Sync cycle exception: {e}")
            
            for _ in range(POLL_INTERVAL_SECONDS):
                if not self.is_running:
                    break
                time.sleep(1)

    def start(self):
        import pystray
        from pystray import MenuItem as item

        if not self.token:
            self.token = prompt_login_dialog()

        sync_thread = threading.Thread(target=self.background_poller, daemon=True)
        sync_thread.start()

        def on_sync_now(icon, item):
            threading.Thread(target=self.run_sync_cycle, daemon=True).start()

        def on_open_excel(icon, item):
            ensure_excel_workbook(self.excel_path)
            open_file(self.excel_path)

        def on_logout(icon, item):
            remove_stored_token()
            self.token = None
            if self.tray_icon:
                self.tray_icon.notify("Logged out from NAMI desktop sync", "NAMI")

        def on_exit(icon, item):
            self.is_running = False
            icon.stop()

        menu = pystray.Menu(
            item('NAMI Desktop Sync (Running)', lambda: None, enabled=False),
            item('Sync Now', on_sync_now),
            item('Open Excel Sheet', on_open_excel),
            item('Logout', on_logout),
            item('Exit', on_exit)
        )

        self.tray_icon = pystray.Icon(
            "nami_sync",
            create_tray_icon(),
            "NAMI Sync Agent",
            menu
        )

        print("NAMI Desktop Agent running in system tray...")
        self.tray_icon.run()

if __name__ == '__main__':
    agent = NamiSyncAgent()
    agent.start()
